"""
Document Generator Module
=========================
Exports meeting minutes to formatted .docx using python-docx.
"""

from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    DOCX_AVAILABLE = True
except Exception:
    # Minimal fallback implementations for environments without python-docx (used in tests)
    DOCX_AVAILABLE = False

    class Document:
        def __init__(self):
            self._paragraphs = []
            self.sections = []

            # Minimal styles container to mimic python-docx for tests
            class DummyStyle:
                def __init__(self):
                    self.font = type("F", (), {"name": None, "size": None})

                    class RFonts:
                        def set(self, *args, **kwargs):
                            pass

                    class RPr:
                        def __init__(self):
                            self.rFonts = RFonts()

                    class Element:
                        def __init__(self):
                            self.rPr = RPr()

                    self._element = Element()

            class Styles:
                def __init__(self):
                    self._styles = {"Normal": DummyStyle()}

                def __getitem__(self, key):
                    return self._styles.setdefault(key, DummyStyle())

            self.styles = Styles()

        class Run:
            def __init__(self, text=""):
                self.text = str(text)
                self.bold = False
                self.italic = False
                self.font = type("F", (), {"size": None, "color": type("C", (), {"rgb": None})()})

        class Paragraph:
            def __init__(self, text=""):
                self.runs = []
                self.paragraph_format = type("PF", (), {"space_after": None})
                self.alignment = None
                if text:
                    self.add_run(text)

            def add_run(self, text=""):
                # Create a lightweight run-like object for fallback
                run = type(
                    "Run",
                    (),
                    {
                        "text": str(text),
                        "bold": False,
                        "italic": False,
                        "font": type(
                            "F", (), {"size": None, "color": type("C", (), {"rgb": None})()}
                        )(),
                    },
                )()
                self.runs.append(run)
                return run

        def add_paragraph(self, text="", **kwargs):
            # Accept style and other kwargs for compatibility
            para = self.Paragraph(text)
            self._paragraphs.append(para)
            return para

        def add_heading(self, text, level=None, **kwargs):
            para = self.Paragraph(text)
            self._paragraphs.append(para)
            return para

        def add_table(self, rows, cols):
            outer = self

            class Cell:
                def __init__(self):
                    self.paragraphs = [outer.Paragraph()]

                    # Minimal _tc structure to support shading and other docx operations in fallback
                    class TCPr:
                        def append(self, *args, **kwargs):
                            pass

                    class TC:
                        def get_or_add_tcPr(self):
                            return TCPr()

                    self._tc = TC()

                @property
                def text(self):
                    if self.paragraphs and self.paragraphs[0].runs:
                        return " ".join(run.text for run in self.paragraphs[0].runs)
                    return ""

                @text.setter
                def text(self, value):
                    # Create lightweight run-like object
                    self.paragraphs[0].runs = [
                        type(
                            "Run",
                            (),
                            {
                                "text": str(value),
                                "bold": False,
                                "italic": False,
                                "font": type(
                                    "F", (), {"size": None, "color": type("C", (), {"rgb": None})()}
                                )(),
                            },
                        )()
                    ]

            class Row:
                def __init__(self, cols):
                    self.cells = [Cell() for _ in range(cols)]

            table = type(
                "Table",
                (),
                {"rows": [Row(cols) for _ in range(rows)], "style": None, "alignment": None},
            )
            return table

        def save(self, path):
            # Save a plain text fallback document so tests can verify file exists
            lines = []
            for p in self._paragraphs:
                if hasattr(p, "runs"):
                    lines.append(" ".join(getattr(r, "text", "") for r in p.runs))
                else:
                    lines.append(str(p))
            with open(path, "w", encoding="utf-8") as f:
                f.write("\n".join(lines))

    class Pt:
        def __init__(self, value):
            self.value = value

    class Cm:
        def __init__(self, value):
            self.value = value

    class RGBColor:
        def __init__(self, r, g, b):
            pass

    class WD_ALIGN_PARAGRAPH:
        CENTER = 1

    class WD_TABLE_ALIGNMENT:
        LEFT = 1

    class OxmlElement:
        def __init__(self, *args, **kwargs):
            pass

        def set(self, *args, **kwargs):
            pass

    def qn(x):
        return x


from src.summarizer import MeetingSummary
from src.transcriber import TranscriptSegment


@dataclass
class MeetingMetadata:
    """Meeting information for document header"""

    title: str
    date: str
    time: str = ""
    location: str = ""
    duration: str = ""
    participants: Optional[List[str]] = None
    organizer: str = ""
    agenda: str = ""

    @classmethod
    def create_default(cls, audio_duration_sec: float = 0) -> "MeetingMetadata":
        """Create default metadata"""
        duration_str = ""
        if audio_duration_sec > 0:
            hours = int(audio_duration_sec // 3600)
            minutes = int((audio_duration_sec % 3600) // 60)
            seconds = int(audio_duration_sec % 60)

            if hours > 0:
                duration_str = f"{hours} jam {minutes} menit {seconds} detik"
            else:
                duration_str = f"{minutes} menit {seconds} detik"

        return cls(
            title="Notulensi Rapat",
            date=datetime.now().strftime("%d %B %Y"),
            time=datetime.now().strftime("%H:%M"),
            duration=duration_str,
        )


@dataclass
class DocumentConfig:
    """Configuration for document generation"""

    # Font settings
    title_font_size: int = 18
    heading1_font_size: int = 14
    heading2_font_size: int = 12
    body_font_size: int = 11
    font_family: str = "Calibri"

    # Layout
    page_width: float = 21.0  # cm (A4)
    page_height: float = 29.7  # cm (A4)
    margin_top: float = 2.5
    margin_bottom: float = 2.5
    margin_left: float = 3.0
    margin_right: float = 2.5

    # Content options
    include_timestamps: bool = True
    include_speaker_colors: bool = True
    include_table_of_contents: bool = False
    include_page_numbers: bool = True

    # Sections to include
    sections: Dict[str, bool] = field(
        default_factory=lambda: {
            "header": True,
            "meeting_info": True,
            "summary": True,
            "decisions": True,
            "action_items": True,
            "transcript": True,
            "footer": True,
        }
    )


class DocumentGenerator:
    """
    Generates formatted .docx meeting minutes.

    Structure:
        - Title
        - Meeting Information
        - Executive Summary
        - Key Points
        - Decisions
        - Action Items
        - Full Transcript
        - Footer

    Attributes:
        config: DocumentConfig object
        output_dir: Output directory path

    Example:
        >>> generator = DocumentGenerator()
        >>> doc_path = generator.generate(metadata, summary, transcript)
        >>> print(f"Document saved: {doc_path}")
    """

    # Speaker colors for visual distinction
    SPEAKER_COLORS = [
        RGBColor(0, 102, 204),  # Blue
        RGBColor(204, 51, 0),  # Red
        RGBColor(0, 153, 51),  # Green
        RGBColor(153, 51, 153),  # Purple
        RGBColor(204, 102, 0),  # Orange
        RGBColor(0, 153, 153),  # Teal
        RGBColor(102, 102, 0),  # Olive
        RGBColor(153, 0, 76),  # Maroon
    ]

    def __init__(self, config: Optional[DocumentConfig] = None, output_dir: str = "./data/output"):
        """
        Initialize DocumentGenerator.

        Args:
            config: DocumentConfig object
            output_dir: Directory for output files
        """
        self.config = config or DocumentConfig()
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self._speaker_color_map: Dict[str, RGBColor] = {}

    def generate(
        self,
        metadata: MeetingMetadata,
        summary: MeetingSummary,
        transcript: List[TranscriptSegment],
        output_filename: Optional[str] = None,
    ) -> str:
        """
        Generate complete meeting minutes document.

        Args:
            metadata: Meeting information
            summary: Generated summary
            transcript: Transcribed segments with speakers
            output_filename: Output file name (auto-generated if None)

        Returns:
            Path to generated document
        """
        # Create document
        doc = Document()

        # Setup document
        self._setup_document(doc)
        self._setup_styles(doc)

        # Build speaker color map
        self._build_speaker_color_map(transcript)

        # Add sections
        if self.config.sections.get("header", True):
            self._add_title(doc, metadata)

        if self.config.sections.get("meeting_info", True):
            self._add_meeting_info(doc, metadata)

        if self.config.sections.get("summary", True):
            self._add_summary_section(doc, summary)

        if self.config.sections.get("decisions", True):
            self._add_decisions_section(doc, summary.decisions)

        if self.config.sections.get("action_items", True):
            self._add_action_items_section(doc, summary.action_items)

        if self.config.sections.get("transcript", True):
            self._add_transcript_section(doc, transcript)

        if self.config.sections.get("footer", True):
            self._add_footer(doc)

        # Generate filename if not provided
        if output_filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_title = self._sanitize_filename(metadata.title)[:30]
            ext = ".docx" if DOCX_AVAILABLE else ".txt"
            output_filename = f"notulensi_{safe_title}_{timestamp}{ext}"

        # Ensure .docx extension
        if not output_filename.endswith(".docx"):
            output_filename = Path(output_filename).with_suffix(".docx").name

        output_path = self.output_dir / output_filename

        # Save document
        if DOCX_AVAILABLE:
            doc.save(str(output_path))
        else:
            # If python-docx is not available, build a minimal valid .docx package so Word can open it.
            warnings.warn(
                "python-docx is not available in the current environment; generating a minimal .docx package instead."
            )
            paragraphs = self._extract_paragraph_texts(doc)
            self._save_minimal_docx(str(output_path), paragraphs)

        return str(output_path)

    def _setup_document(self, doc: Document):
        """Configure document settings"""
        # Set page margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Cm(self.config.margin_top)
            section.bottom_margin = Cm(self.config.margin_bottom)
            section.left_margin = Cm(self.config.margin_left)
            section.right_margin = Cm(self.config.margin_right)

    def _setup_styles(self, doc: Document):
        """Configure document styles"""
        # Normal style
        style = doc.styles["Normal"]
        style.font.name = self.config.font_family
        style.font.size = Pt(self.config.body_font_size)

        # Set font for East Asian text
        style._element.rPr.rFonts.set(qn("w:eastAsia"), self.config.font_family)

    def _build_speaker_color_map(self, transcript: List[TranscriptSegment]):
        """Build consistent color mapping for speakers"""
        speakers = sorted(set(seg.speaker_id for seg in transcript))

        for i, speaker in enumerate(speakers):
            self._speaker_color_map[speaker] = self.SPEAKER_COLORS[i % len(self.SPEAKER_COLORS)]

    def _add_title(self, doc: Document, metadata: MeetingMetadata):
        """Add document title"""
        # Main title
        title_para = doc.add_paragraph()
        title_run = title_para.add_run("NOTULENSI RAPAT")
        title_run.bold = True
        title_run.font.size = Pt(self.config.title_font_size)
        title_run.font.color.rgb = RGBColor(0, 51, 102)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subtitle with meeting title
        if metadata.title and metadata.title != "Notulensi Rapat":
            subtitle_para = doc.add_paragraph()
            subtitle_run = subtitle_para.add_run(metadata.title)
            subtitle_run.bold = True
            subtitle_run.font.size = Pt(self.config.heading1_font_size)
            subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Generated by note
        note_para = doc.add_paragraph()
        note_run = note_para.add_run("Generated by AI Meeting Transcriber (SpeechBrain + BERT)")
        note_run.italic = True
        note_run.font.size = Pt(9)
        note_run.font.color.rgb = RGBColor(128, 128, 128)
        note_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Spacer
        doc.add_paragraph()

    def _add_meeting_info(self, doc: Document, metadata: MeetingMetadata):
        """Add meeting information section"""
        # Section heading
        heading = doc.add_heading("Informasi Rapat", level=1)
        heading.runs[0].font.size = Pt(self.config.heading1_font_size)

        # Create info table
        info_items = [
            ("Tanggal", metadata.date),
            ("Waktu", metadata.time or "-"),
            ("Lokasi/Platform", metadata.location or "-"),
            ("Durasi", metadata.duration or "-"),
            ("Penyelenggara", metadata.organizer or "-"),
        ]

        # Filter out empty items
        info_items = [(label, value) for label, value in info_items if value and value != "-"]

        if info_items:
            table = doc.add_table(rows=len(info_items), cols=2)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT

            for i, (label, value) in enumerate(info_items):
                row = table.rows[i]

                # Label cell
                cell_label = row.cells[0]
                cell_label.text = label
                cell_label.paragraphs[0].runs[0].bold = True
                cell_label.width = Cm(4)

                # Value cell
                cell_value = row.cells[1]
                cell_value.text = value

        # Add participants if available
        if metadata.participants:
            doc.add_paragraph()
            para = doc.add_paragraph()
            para.add_run("Peserta Rapat: ").bold = True
            para.add_run(", ".join(metadata.participants))

        # Add agenda if available
        if metadata.agenda:
            doc.add_paragraph()
            para = doc.add_paragraph()
            para.add_run("Agenda: ").bold = True
            para.add_run(metadata.agenda)

        # Spacer
        doc.add_paragraph()

    def _add_summary_section(self, doc: Document, summary: MeetingSummary):
        """Add executive summary section"""
        # Section heading
        heading = doc.add_heading("Ringkasan Eksekutif", level=1)
        heading.runs[0].font.size = Pt(self.config.heading1_font_size)

        # Overview
        if summary.overview and not self._is_placeholder_text(summary.overview):
            overview_para = doc.add_paragraph()
            overview_para.add_run(summary.overview)
            overview_para.paragraph_format.space_after = Pt(12)
        else:
            overview_para = doc.add_paragraph()
            overview_para.add_run(
                "Ringkasan tidak tersedia. (Model ringkasan tidak dimuat atau data tidak mencukupi.)"
            )
            overview_para.runs[0].italic = True
            overview_para.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        # Key points (filter placeholders)
        filtered_points = [
            p for p in (summary.key_points or []) if not self._is_placeholder_text(p)
        ]
        if filtered_points:
            subheading = doc.add_heading("Poin-Poin Penting", level=2)
            subheading.runs[0].font.size = Pt(self.config.heading2_font_size)

            for point in filtered_points:
                para = doc.add_paragraph(point, style="List Bullet")
        else:
            para = doc.add_paragraph()
            para.add_run("Tidak ada poin penting yang dihasilkan secara otomatis.")
            para.runs[0].italic = True
            para.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        # Topics discussed (filter placeholders)
        topics_filtered = [t for t in (summary.topics or []) if not self._is_placeholder_text(t)]
        if topics_filtered:
            doc.add_paragraph()
            para = doc.add_paragraph()
            para.add_run("Topik yang dibahas: ").bold = True
            para.add_run(", ".join(topics_filtered))

        # Spacer
        doc.add_paragraph()

    def _add_decisions_section(self, doc: Document, decisions: List[str]):
        """Add decisions section"""
        # Section heading
        heading = doc.add_heading("Keputusan Rapat", level=1)
        heading.runs[0].font.size = Pt(self.config.heading1_font_size)

        if decisions:
            for i, decision in enumerate(decisions, 1):
                para = doc.add_paragraph()
                para.add_run(f"{i}. ").bold = True
                para.add_run(decision)
        else:
            para = doc.add_paragraph()
            para.add_run("Tidak ada keputusan yang teridentifikasi secara otomatis.")
            para.runs[0].italic = True
            para.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        # Spacer
        doc.add_paragraph()

    def _add_action_items_section(self, doc: Document, action_items: List[Dict[str, str]]):
        """Add action items section"""
        # Section heading
        heading = doc.add_heading("Action Items / Tindak Lanjut", level=1)
        heading.runs[0].font.size = Pt(self.config.heading1_font_size)

        if action_items:
            # Create table
            table = doc.add_table(rows=len(action_items) + 1, cols=4)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT

            # Header row
            headers = ["No.", "Penanggung Jawab", "Tugas", "Deadline"]
            header_row = table.rows[0]

            for i, header_text in enumerate(headers):
                cell = header_row.cells[i]
                cell.text = header_text

                # Style header
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

                # Set header background color
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E2F3")
                cell._tc.get_or_add_tcPr().append(shading)

            # Data rows
            for i, item in enumerate(action_items, 1):
                row = table.rows[i]

                row.cells[0].text = str(i)
                row.cells[1].text = item.get("owner", "-")
                row.cells[2].text = item.get("task", "-")
                row.cells[3].text = item.get("due", "-")

            # Set column widths
            for row in table.rows:
                row.cells[0].width = Cm(1.0)
                row.cells[1].width = Cm(3.5)
                row.cells[2].width = Cm(9.0)
                row.cells[3].width = Cm(2.5)
        else:
            para = doc.add_paragraph()
            para.add_run("Tidak ada action item yang teridentifikasi secara otomatis.")
            para.runs[0].italic = True
            para.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        # Spacer
        doc.add_paragraph()

    def _add_transcript_section(self, doc: Document, transcript: List[TranscriptSegment]):
        """Add full transcript section"""
        # Section heading
        heading = doc.add_heading("Transkrip Percakapan", level=1)
        heading.runs[0].font.size = Pt(self.config.heading1_font_size)

        if not transcript:
            para = doc.add_paragraph()
            para.add_run("Tidak ada transkrip yang tersedia.")
            para.runs[0].italic = True
            return

        # Add each segment
        for seg in transcript:
            para = doc.add_paragraph()

            # Timestamp
            if self.config.include_timestamps:
                timestamp = self._format_timestamp(seg.start, seg.end)

                # Speaker label with color
                speaker_run = para.add_run(f"{seg.speaker_id} [{timestamp}]: ")
                speaker_run.bold = True

                if self.config.include_speaker_colors:
                    color = self._speaker_color_map.get(seg.speaker_id, RGBColor(0, 0, 0))
                    speaker_run.font.color.rgb = color
            else:
                speaker_run = para.add_run(f"{seg.speaker_id}: ")
                speaker_run.bold = True

            # Transcript text (sanitize placeholder/fallback strings)
            text = seg.text or ""
            cleaned = self._clean_text_for_doc(text)
            para.add_run(cleaned)

            # Mark overlapping speech
            if seg.is_overlap:
                overlap_run = para.add_run(" [OVERLAP]")
                overlap_run.italic = True
                overlap_run.font.color.rgb = RGBColor(255, 102, 0)
                overlap_run.font.size = Pt(9)

    def _add_footer(self, doc: Document):
        """Add document footer"""
        # Separator line
        doc.add_paragraph()
        separator = doc.add_paragraph("─" * 70)
        separator.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Footer text
        footer_para = doc.add_paragraph()

        timestamp = datetime.now().strftime("%d %B %Y, %H:%M:%S")
        footer_text = f"Dokumen ini dihasilkan secara otomatis pada {timestamp}"

        footer_run = footer_para.add_run(footer_text)
        footer_run.italic = True
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor(128, 128, 128)
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Disclaimer
        disclaimer_para = doc.add_paragraph()
        disclaimer_text = (
            "Hasil transkripsi dan ringkasan mungkin mengandung ketidakakuratan. "
            "Harap verifikasi informasi penting."
        )

        disclaimer_run = disclaimer_para.add_run(disclaimer_text)
        disclaimer_run.italic = True
        disclaimer_run.font.size = Pt(8)
        disclaimer_run.font.color.rgb = RGBColor(150, 150, 150)
        disclaimer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _is_placeholder_text(self, text: Optional[str]) -> bool:
        """Detect summarizer/ASR fallback placeholder text."""
        if not text:
            return True
        t = str(text).strip()
        # common placeholder patterns from summarizer / transcriber fallbacks
        if re.search(r"\[\s*Transkripsi placeholder", t, re.I):
            return True
        if re.search(r"placeholder", t, re.I) and len(t) < 120:
            return True
        return False

    def _clean_text_for_doc(self, text: Optional[str]) -> str:
        """Clean text for document: replace raw placeholders with user-friendly notices."""
        if not text or self._is_placeholder_text(text):
            return "[transkripsi tidak tersedia]"
        # Remove any bracketed placeholder fragments embedded in text
        cleaned = re.sub(r"\[\s*Transkripsi placeholder[^\]]*\]", "", str(text), flags=re.I).strip()
        return cleaned or "[transkripsi tidak tersedia]"

    @staticmethod
    def _format_timestamp(start: float, end: float) -> str:
        """Format time range as HH:MM:SS"""

        def sec_to_str(sec: float) -> str:
            sec = max(0.0, float(sec))
            h = int(sec // 3600)
            m = int((sec % 3600) // 60)
            s = int(sec % 60)

            if h > 0:
                return f"{h:02d}:{m:02d}:{s:02d}"
            return f"{m:02d}:{s:02d}"

        return f"{sec_to_str(start)}–{sec_to_str(end)}"

    def _save_minimal_docx(self, path: str, paragraphs: List[str]):
        """Create a minimal valid .docx (zip package) containing plain paragraphs.
        This is a lightweight fallback when python-docx is not installed, to ensure
        the generated file can be opened in Word.
        """
        import zipfile

        def _escape_xml(s: str) -> str:
            return (
                s.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace('"', "&quot;")
                .replace("'", "&apos;")
            )

        content_types = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n'
            '  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\n'
            '  <Default Extension="xml" ContentType="application/xml"/>\n'
            '  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>\n'
            "</Types>"
        )

        rels = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>\n'
            "</Relationships>"
        )

        doc_xml_header = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:w10="urn:schemas-microsoft-com:office:word" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
            'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
            'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
            'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">\n'
            "  <w:body>\n"
        )

        doc_xml_footer = (
            "    <w:sectPr>\n"
            '      <w:pgSz w:w="11900" w:h="16840"/>\n'
            '      <w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/>\n'
            "    </w:sectPr>\n"
            "  </w:body>\n"
            "</w:document>"
        )

        # Build paragraphs as simple <w:p><w:r><w:t>text</w:t></w:r></w:p>
        paras_xml = []
        for p in paragraphs:
            t = _escape_xml(p.strip())
            if not t:
                # preserve blank line
                paras_xml.append("    <w:p/>\n")
            else:
                paras_xml.append(f'    <w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>\n')

        doc_xml = doc_xml_header + "".join(paras_xml) + doc_xml_footer

        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", content_types)
            z.writestr("_rels/.rels", rels)
            z.writestr("word/document.xml", doc_xml)

    def _extract_paragraph_texts(self, doc: Document) -> List[str]:
        """Get paragraphs text for python-docx Document or fallback Document"""
        paras: List[str] = []
        # python-docx Document
        try:
            # using attribute if present
            if hasattr(doc, "paragraphs"):
                for p in doc.paragraphs:
                    paras.append(p.text)
                return paras
        except Exception:
            pass

        # fallback minimal Document implementation
        if hasattr(doc, "_paragraphs"):
            for p in doc._paragraphs:
                if hasattr(p, "runs"):
                    paras.append(" ".join(getattr(r, "text", "") for r in p.runs))
                else:
                    paras.append(str(p))
        return paras

    @staticmethod
    def _sanitize_filename(filename: str) -> str:
        """Remove invalid characters from filename"""
        import re

        # Remove invalid characters
        sanitized = re.sub(r'[<>:"/\\|?*]', "", filename)
        # Replace spaces with underscores
        sanitized = sanitized.replace(" ", "_")
        return sanitized
